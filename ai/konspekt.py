"""
ai/konspekt.py — "Konspekt qilish" xizmati: uzun matn yoki hujjatni (.docx/.pdf)
qisqa, TUZILGAN konspektga (ma'ruza xulosasiga) aylantiradi — imtihonga
tayyorgarlik uchun asosiy fikrlar, ta'riflar, sana/raqamlar saqlanadi.

ai/davom_yoz.py dan farqi: bu yerda matn DAVOM ETTIRILMAYDI, aksincha
QISQARTIRILADI va bo'limlarga ajratiladi.
"""

import logging

from docx import Document
from pypdf import PdfReader

from config import KONSPEKT_BASE_URL, KONSPEKT_KEY, KONSPEKT_MODEL
from ai._runtime import client_va_model

log = logging.getLogger("konspekt")

SOZ_LIMIT = 20000

_SYSTEM = (
    "Sen professional o'qituvchi-muharrirsan. Senga uzun matn yoki ma'ruza matni "
    "beriladi. Vazifang — uni IMTIHONGA TAYYORGARLIK uchun qisqa, tuzilgan "
    "KONSPEKTGA (xulosaga) aylantirish.\n\n"
    "QAT'IY QOIDALAR:\n"
    "- Faqat MUHIM narsalarni saqla: ta'riflar, asosiy faktlar, raqamlar, "
    "sanalar, formulalar, tushunchalar — ortiqcha so'z, misol yoki tavsifni "
    "OLIB TASHLA.\n"
    "- Matnni MAZMUNIY BO'LIMLARGA AJRAT — har bo'lim uchun qisqa sarlavha yoz.\n"
    "- Har bo'lim ostida QISQA bandlar bilan yoz (har band — bitta aniq fikr).\n"
    "- LOTIN alifbosida (kirill EMAS), aniq va ravon tilda yoz.\n"
    "- Original matnda YO'Q narsani O'YLAB TOPMA — faqat matnda bor faktlarni ishlat.\n"
    "- Quyidagi formatda yoz (Markdown emas, oddiy matn):\n\n"
    "📌 <asosiy mavzu nomi>\n\n"
    "🔹 <bo'lim sarlavhasi>\n"
    "— band\n— band\n\n"
    "🔹 <keyingi bo'lim sarlavhasi>\n— band\n...\n\n"
    "- Kirish gap ('Mana konspekt:' kabi) yoki izoh YOZMA — to'g'ridan-to'g'ri "
    "konspektning o'zini qaytar."
)


def _docxdan_matn_ol(yol):
    doc = Document(yol)
    qismlar = [p.text for p in doc.paragraphs if p.text.strip()]
    for jadval in doc.tables:
        for qator in jadval.rows:
            for katak in qator.cells:
                qismlar += [p.text for p in katak.paragraphs if p.text.strip()]
    return "\n".join(qismlar)


def _pdfdan_matn_ol(yol):
    reader = PdfReader(yol)
    sahifalar = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(s for s in sahifalar if s.strip())


def fayldan_matn_ol(yol):
    """.docx yoki .pdf fayldan matn oladi (kengaytmaga qarab)."""
    if yol.lower().endswith(".pdf"):
        return _pdfdan_matn_ol(yol)
    return _docxdan_matn_ol(yol)


def konspekt_qil(matn: str) -> str:
    """Matnni qisqa, tuzilgan konspektga aylantiradi.
    Matn bo'sh yoki juda katta bo'lsa, yoki AI konspekt yasay olmasa — ValueError."""
    matn = (matn or "").strip()
    if not matn:
        raise ValueError("matn topilmadi (bo'sh yoki skanerlangan rasm bo'lishi mumkin)")
    soz = len(matn.split())
    if soz > SOZ_LIMIT:
        raise ValueError(f"matn juda katta ({soz} so'z, limit {SOZ_LIMIT} so'z)")

    try:
        client, model = client_va_model("konspekt", KONSPEKT_BASE_URL, KONSPEKT_KEY, KONSPEKT_MODEL)
        resp = client.chat.completions.create(
            model=model, temperature=0.3,
            messages=[{"role": "system", "content": _SYSTEM},
                      {"role": "user", "content": matn}])
        natija = (resp.choices[0].message.content or "").strip()
    except Exception as e:
        log.error("Konspekt qilishda xato: %s", e)
        raise

    if not natija:
        raise ValueError("AI konspekt yasay olmadi, qaytadan urinib ko'ring")
    return natija
