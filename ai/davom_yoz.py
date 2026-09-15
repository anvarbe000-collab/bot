"""
ai/davom_yoz.py — "Davom ettirish" xizmati: chala qolgan yoki to'liq bo'lmagan
matnni (.docx yoki .pdf) o'qib, AI yordamida xuddi shu uslub va mavzuda DAVOM
ETTIRIB, mukammal va TO'LIQ hujjat (.docx) yasaydi.

ai/docx_tuzat.py dan farqi: bu yerda format SAQLASH shart emas — original matn
faqat MATN sifatida o'qiladi (docx yoki pdf'dan), natija esa yangi, toza
formatlangan .docx fayl (asl matn + davomi, bitta yaxlit hujjat).
"""

import logging

from docx import Document
from pypdf import PdfReader
from openai import OpenAI

from config import TUZAT_BASE_URL, TUZAT_KEY, TUZAT_MODEL

log = logging.getLogger("davom-yoz")

_client = OpenAI(base_url=TUZAT_BASE_URL, api_key=TUZAT_KEY)

SOZ_LIMIT = 15000

_SYSTEM = (
    "Sen professional o'zbek tilida (LOTIN alifbosida) yozuvchi-muharrirsan. "
    "Senga foydalanuvchining CHALA QOLGAN yoki TO'LIQ BO'LMAGAN matni beriladi. "
    "Vazifang: matnni diqqat bilan o'qib, uning MAVZUSI, USLUBI, OHANGI va "
    "MANTIQIY YO'NALISHINI saqlagan holda DAVOM ETTIRIB, uni to'liq, mukammal "
    "va tugallangan holga keltirish.\n\n"
    "QAT'IY QOIDALAR:\n"
    "- FAQAT matnning DAVOMINI yoz — original matnni QAYTA YOZMA yoki takrorlama.\n"
    "- LOTIN alifbosida, adabiy va ravon tilda yoz.\n"
    "- Original bilan bir xil uslub, ohang va shaxsda (1-shaxs/3-shaxs, rasmiy/"
    "norasmiy) davom et — birdan uslub o'zgarib qolmasin.\n"
    "- Matn qanday yakunlanishi kerak bo'lsa (xulosa, yakuniy fikr, natija) — "
    "shunday tabiiy tarzda yakunla, chala qoldirma.\n"
    "- Izoh, tushuntirish, sarlavha qo'shma — faqat davom matnini qaytar."
)


def _docxdan_matn_ol(yol):
    doc = Document(yol)
    qismlar = [p.text for p in doc.paragraphs if p.text.strip()]
    for jadval in doc.tables:
        for qator in jadval.rows:
            for katak in qator.cells:
                qismlar += [p.text for p in katak.paragraphs if p.text.strip()]
    return "\n".join(qismlar)


def _pdfdan_matn_ol(yol):
    reader = PdfReader(yol)
    sahifalar = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(s for s in sahifalar if s.strip())


def _fayldan_matn_ol(yol):
    if yol.lower().endswith(".pdf"):
        return _pdfdan_matn_ol(yol)
    return _docxdan_matn_ol(yol)


def davom_ettir(kirish_yol: str, chiqish_yol: str, soz_soni: int = None):
    """Faylni (.docx yoki .pdf) o'qib, matnni AI bilan davom ettiradi va
    yangi .docx faylga (asl matn + davomi) saqlaydi.
    soz_soni berilsa, davom qism taxminan shuncha so'zdan iborat bo'lishga harakat qiladi.

    Qaytaradi: (chiqish_yol, asl_soz_soni, davom_soz_soni)
    Hujjat bo'sh yoki juda katta bo'lsa, yoki AI davom yasay olmasa — ValueError.
    """
    asl_matn = _fayldan_matn_ol(kirish_yol).strip()
    if not asl_matn:
        raise ValueError("hujjatdan matn topilmadi (bo'sh yoki skanerlangan rasm bo'lishi mumkin)")

    asl_soz = len(asl_matn.split())
    if asl_soz > SOZ_LIMIT:
        raise ValueError(f"hujjat juda katta ({asl_soz} so'z, limit {SOZ_LIMIT} so'z)")

    system = _SYSTEM
    if soz_soni:
        system += (
            f"\n\nDAVOM ETTIRILGAN QISM TAXMINAN {soz_soni} SO'ZDAN iborat bo'lsin "
            f"(bir oz kam yoki ko'p bo'lishi mumkin, lekin shu miqdorga yaqin bo'lsin)."
        )

    try:
        resp = _client.chat.completions.create(
            model=TUZAT_MODEL, temperature=0.6,
            messages=[{"role": "system", "content": system},
                      {"role": "user", "content": asl_matn}])
        davomi = (resp.choices[0].message.content or "").strip()
    except Exception as e:
        log.error("Davom ettirishda xato: %s", e)
        raise

    if not davomi:
        raise ValueError("AI matn davomini yasay olmadi, qaytadan urinib ko'ring")

    doc = Document()
    for qator in asl_matn.split("\n"):
        if qator.strip():
            doc.add_paragraph(qator.strip())
    for qator in davomi.split("\n"):
        if qator.strip():
            doc.add_paragraph(qator.strip())
    doc.save(chiqish_yol)

    return chiqish_yol, asl_soz, len(davomi.split())
